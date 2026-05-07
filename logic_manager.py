import os
import json
import re
import subprocess
import sys
import shutil
from pathlib import Path
from urllib.parse import urlparse
from playwright.sync_api import sync_playwright
from core.setup_runner import SetupRunner
from core.ai_generator import AIGenerator

class AutomationLogic:
    def __init__(self):
        self.base_dir = Path(__file__).resolve().parent
        self.ai = AIGenerator()
        self._vector_db = None
        self._ai_config = None

    def set_ai_components(self, vector_db=None, ai_config=None):
        """Truyen VectorDB va AIConfig vao logic manager."""
        if vector_db is not None:
            self._vector_db = vector_db
        if ai_config is not None:
            self._ai_config = ai_config
        self.ai.vector_db = self._vector_db
        self.ai.ai_config = self._ai_config

    def get_site_folder_name(self, url):
        """Trích xuất tên thư mục từ domain của URL"""
        parsed_url = urlparse(url)
        domain = parsed_url.netloc.replace('.', '_') if parsed_url.netloc else "unknown_site"
        return re.sub(r'[\\/*?:"<>|]', "", domain)

    def get_setup_scripts(self, project_path):
        """List available setup scripts in scripts/setup/."""
        return SetupRunner.list_scripts(project_path)

    def scan_url(self, url, browser_name="chromium", setup_script=""):
        """Thực hiện quét trang web bằng Playwright và trả về danh sách phần tử bao gồm cả các thẻ rỗng

        Args:
            url: Target URL to scan.
            browser_name: Browser to use (chromium/firefox/webkit).
            setup_script: Optional absolute path to a setup script to run before scanning.
        """
        if not url:
            raise ValueError("URL không được để trống")

        with sync_playwright() as p:
            browser_type = getattr(p, browser_name, p.chromium)
            browser = browser_type.launch(headless=True)
            page = browser.new_page()

            # Run setup script if provided (e.g. login, dismiss banners)
            if setup_script:
                SetupRunner.run(page, setup_script)

            page.goto(url, timeout=60000)
            
            # Trong file logic_manager.txt, thay đổi đoạn evaluate trong scan_url:

            elements = page.evaluate(r"""
                () => {
                    const results = [];
                    const selectors = 'body, div, section, input, button, select, textarea, h1, h2, h3, h4, h5, h6, span, p, a, label';
                    
                    // Hàm phụ để lấy selector rút gọn của một phần tử
                    const getSimpleSelector = (el) => {
                        if (el.id) return `#${el.id}`;
                        if (el.getAttribute('data-test')) return `[data-test="${el.getAttribute('data-test')}"]`;
                        if (el.className && typeof el.className === 'string' && el.className.trim() !== '') {
                            return '.' + el.className.trim().split(/\s+/)[0];
                        }
                        return el.tagName.toLowerCase();
                    };

                    document.querySelectorAll(selectors).forEach(el => {
                        let tagName = el.tagName.toLowerCase();
                        if (['script', 'style', 'noscript', 'meta', 'link'].includes(tagName)) return;

                        let selector = '';
                        const dataTest = el.getAttribute('data-test');
                        const nameAttr = el.getAttribute('name');                                 

                        // 1. Ưu tiên ID
                        if (el.id) {
                            selector = `#${el.id}`;
                        } 
                        // 2. Ưu tiên data-test
                        else if (dataTest) {
                            selector = `[data-test="${dataTest}"]`;
                        } 
                        // 2. Ưu tiên name
                        else if (nameAttr) {
                            selector = `[name="${nameAttr}"]`;
                        }
                        else {
                            // Lấy thông tin 2 cấp cha
                            let p1 = el.parentElement;
                            let p2 = p1 ? p1.parentElement : null;
                            let prefix = "";
                            if (p2) prefix += getSimpleSelector(p2) + " > ";
                            if (p1) prefix += getSimpleSelector(p1) + " > ";

                            // 3. Ưu tiên Class (có 2 cấp cha)
                            if (el.className && typeof el.className === 'string' && el.className.trim() !== '') {
                                const classes = el.className.trim().split(/\s+/).filter(c => c && !c.includes(':')).join('.');
                                selector = prefix + (classes ? `.${classes}` : tagName);
                            } 
                            // 4. Ưu tiên Thẻ (có 2 cấp cha)
                            else {
                                selector = prefix + tagName;
                            }
                        }

                        // ... (giữ nguyên phần logic xử lý nameCandidate, action, is_info bên dưới) [cite: 9, 10, 12, 14, 15]
                        let innerTxt = el.innerText ? el.innerText.split('\n')[0].trim().substring(0, 20) : '';
                        let nameCandidate = el.id || dataTest || el.getAttribute('name') || innerTxt || tagName;
                        let name = nameCandidate.toUpperCase().split(/[^A-Z0-9]+/).filter(x => x).join('_');
                        if (!name) name = tagName.toUpperCase() + '_ELEMENT';

                        let action = 'click';
                        let is_info = false;
                        if (['input', 'textarea'].includes(tagName)) {
                            action = 'fill';
                        } else if (tagName === 'select') {
                            action = 'select';
                        } else {
                            if (!['button', 'a'].includes(tagName)) {
                                action = 'verify_text';
                                is_info = true;
                            }
                        }
                        if (el.type === 'password') action = 'fill_password';

                        if (selector && selector !== '.') {
                            results.push({ 
                                name: name.substring(0, 50),
                                type: is_info ? 'Info' : 'Action',
                                selector: selector,
                                action: action
                            });
                        }
                    });
                    // ... (giữ nguyên phần filter trùng lặp) [cite: 19, 20]
                    const seen = new Set();
                    return results.filter(item => {
                        const duplicate = seen.has(item.selector);
                        seen.add(item.selector);
                        return !duplicate;
                    });
                }
            """)
            browser.close()
            return elements

    # def save_configuration(self, project_path, target_url, page_id, data_rows):
        
    #     p_path = Path(project_path)
    #     site_folder_name = self.get_site_folder_name(target_url)

    #     locators = {}
    #     workflow = []
        
    #     for row in data_rows:
    #         name, p_type, selector, action = row
    #         locators[name] = {"selector": selector, "type": p_type.lower()}
            
    #         workflow.append({"id": name, "action": action, "data_key": name.lower()})

    #     # Lưu Locators
    #     loc_site_dir = p_path / "locators" / site_folder_name
    #     loc_site_dir.mkdir(parents=True, exist_ok=True)
    #     with open(loc_site_dir / f"{page_id}.json", "w", encoding="utf-8") as f:
    #         json.dump(locators, f, indent=4, ensure_ascii=False)
            
    #     # Lưu Workflow
    #     tmpl_site_dir = p_path / "templates" / site_folder_name
    #     tmpl_site_dir.mkdir(parents=True, exist_ok=True)
    #     with open(tmpl_site_dir / f"{page_id}_workflow.json", "w", encoding="utf-8") as f:
    #         json.dump({"page_id": page_id, "url": target_url, "steps": workflow}, f, indent=4, ensure_ascii=False)
        
    #     return site_folder_name

        # --- HÀM LƯU JSON MỚI --- #sửa
    def save_locator_json(self, project_path, target_url, page_id, data): #sửa
        """Lưu file locator JSON"""
        p_path = Path(project_path)
        site_folder = self.get_site_folder_name(target_url)
        loc_dir = p_path / "locators" / site_folder
        loc_dir.mkdir(parents=True, exist_ok=True)
        
        with open(loc_dir / f"{page_id}.json", "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)

    def save_template_json(self, project_path, target_url, page_id, data): #sửa
        """Lưu file template JSON"""
        p_path = Path(project_path)
        site_folder = self.get_site_folder_name(target_url)
        tmpl_dir = p_path / "templates" / site_folder
        tmpl_dir.mkdir(parents=True, exist_ok=True)
        
        with open(tmpl_dir / f"{page_id}_workflow.json", "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)

    def import_test_data(self, project_path, target_url, source_file_path):
        """Sao chép file XLSX vào thư mục test_data của dự án"""
        p_path = Path(project_path)
        site_folder = self.get_site_folder_name(target_url)
        dest_dir = p_path / "test_data" / site_folder
        dest_dir.mkdir(parents=True, exist_ok=True)
        
        file_name = Path(source_file_path).name
        dest_path = dest_dir / file_name
        shutil.copy2(source_file_path, dest_path)
        return file_name

    def get_data_files(self, project_path, target_url):
        """Lấy danh sách các file dữ liệu test cho trang web hiện tại"""
        p_path = Path(project_path)
        site_folder = self.get_site_folder_name(target_url)
        data_dir = p_path / "test_data" / site_folder
        
        if not data_dir.exists():
            return []
        
        return [f.name for f in data_dir.glob("*.xlsx")]

    def get_template_files(self, project_path):
        """List all template workflow JSON files across all sites (excludes E2E)."""
        p_path = Path(project_path) / "templates"
        if not p_path.exists():
            return []
        results = []
        for f in sorted(p_path.rglob("*_workflow.json")):
            if f.name.startswith("e2e_"):
                continue
            rel = f.relative_to(p_path)
            results.append(str(rel))
        return results

    def get_e2e_workflow_files(self, project_path):
        """List all E2E workflow JSON files across all sites."""
        p_path = Path(project_path) / "templates"
        if not p_path.exists():
            return []
        results = []
        for f in sorted(p_path.rglob("e2e_*.json")):
            rel = f.relative_to(p_path)
            results.append(str(rel))
        return results

    def get_all_data_files(self, project_path):
        """List all Excel data files across all sites."""
        p_path = Path(project_path) / "test_data"
        if not p_path.exists():
            return []
        results = []
        for f in sorted(p_path.rglob("*.xlsx")):
            rel = f.relative_to(p_path)
            results.append(str(rel))
        return results

    def get_template_info(self, project_path, template_rel_path):
        """Read a template JSON and return its content (page_id, url, steps)."""
        full_path = Path(project_path) / "templates" / template_rel_path
        if not full_path.exists():
            return None
        with open(full_path, "r", encoding="utf-8") as f:
            return json.load(f)

    def merge_templates_to_e2e(self, project_path, template_paths, workflow_id="", description=""):
        """Merge multiple single-page templates into one E2E workflow.

        Args:
            project_path: Root project directory.
            template_paths: List of relative template paths (under templates/).
            workflow_id: Optional ID for the workflow; auto-generated if empty.
            description: Optional description.

        Returns:
            dict with the merged E2E workflow JSON.
        """
        all_pages = []
        all_steps = []
        start_url = ""

        for tpl_rel in template_paths:
            info = self.get_template_info(project_path, tpl_rel)
            if not info:
                continue

            page_id = info.get("page_id", Path(tpl_rel).stem.replace("_workflow", ""))
            url = info.get("url", "")

            if not start_url:
                start_url = url

            if page_id not in all_pages:
                all_pages.append(page_id)

            for step in info.get("steps", []):
                merged_step = dict(step)
                merged_step["page_id"] = page_id
                all_steps.append(merged_step)

        if not workflow_id:
            workflow_id = "_to_".join(all_pages)

        return {
            "workflow_id": workflow_id,
            "type": "e2e",
            "description": description or f"E2E: {' -> '.join(all_pages)}",
            "start_url": start_url,
            "pages": all_pages,
            "steps": all_steps,
        }

    def save_e2e_workflow(self, project_path, site_folder, filename, data):
        """Save an E2E workflow JSON file."""
        p_path = Path(project_path) / "templates" / site_folder
        p_path.mkdir(parents=True, exist_ok=True)
        filepath = p_path / filename
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)
        return str(filepath)

    def get_setup_script_path(self, project_path, script_name):
        """Return absolute path for a setup script by name."""
        if not script_name:
            return ""
        return str(Path(project_path) / "scripts" / "setup" / script_name)

    def get_pytest_command(self, test_file):
        """Tạo lệnh chạy pytest"""
        return [sys.executable, "-m", "pytest", test_file, "-v", "-s", "--tb=no"]

    # --- AI Generation --- #

    def configure_ai(self, api_key):
        """Set the Gemini API key."""
        self.ai.configure(api_key)

    def ai_generate(self, use_case_text, locators, template, page_id, url, project_path=""):
        """Call Gemini AI to generate test artifacts from use case + scanned data."""
        return self.ai.generate(use_case_text, locators, template, page_id, url, project_path)

    def save_test_data_from_rows(self, project_path, target_url, page_id, rows, headers):
        """Save test data rows as an Excel file.

        Args:
            project_path: Root project dir.
            target_url: Target URL (for site folder).
            page_id: Page identifier.
            rows: List of dicts, each dict is a row of test data.
            headers: List of column names.
        """
        import openpyxl
        p_path = Path(project_path)
        site_folder = self.get_site_folder_name(target_url)
        dest_dir = p_path / "test_data" / site_folder
        dest_dir.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{page_id}_data"

        # Write headers
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        # Write data rows
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=row_idx, column=col_idx, value=row_data.get(header, ""))

        filename = f"{page_id}_ai_data.xlsx"
        filepath = dest_dir / filename
        wb.save(filepath)
        return str(filepath)

    def save_setup_script(self, project_path, script_name, data):
        """Save a setup script JSON file."""
        p_path = Path(project_path) / "scripts" / "setup"
        p_path.mkdir(parents=True, exist_ok=True)
        filepath = p_path / script_name
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)
        return str(filepath)